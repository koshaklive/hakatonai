from __future__ import annotations

from io import BytesIO

from docx import Document

from app.core.validator import check_draft_file, validate_kb_file


def _docx_bytes(document: Document) -> bytes:
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_validate_numbered_docx_checks_list_sequence_and_empty_cells():
    document = Document()
    document.add_paragraph("1. РАЗДЕЛ ПЕРВЫЙ")
    document.add_paragraph("1.1 Основной пункт")
    document.add_paragraph("1. первый шаг")
    document.add_paragraph("3. третий шаг")
    document.add_paragraph("1.2 Следующий пункт")
    document.add_paragraph("1.3 Последний пункт")
    document.add_table(rows=1, cols=2).cell(0, 0).text = "Заполнено"

    result = validate_kb_file(_docx_bytes(document), "methodics.docx")

    assert result["kind"] == "numbered"
    assert any("идёт как [1, 3]" in issue["message"] for issue in result["errors"])
    assert any(issue["type"] == "empty_table_cell" for issue in result["errors"])
    assert result["diff"] == {"added": [], "changed": [], "removed": []}


def test_validate_module_checks_note_markers_and_sequence():
    document = Document()
    document.add_paragraph("Шаблон [2]")
    document.add_paragraph("Примечания по заполнению")
    document.add_paragraph("1. Первое примечание")
    document.add_paragraph("3. Третье примечание")
    document.add_paragraph("Общие правила заполнения")

    result = validate_kb_file(_docx_bytes(document), "module.docx")

    assert result["kind"] == "module"
    assert any("маркеры [2]" in issue["message"] for issue in result["errors"])
    assert any("нумерация примечаний не сплошная" in issue["message"] for issue in result["errors"])


def test_check_draft_reports_placeholders_and_manual_semantic_check():
    document = Document()
    document.add_paragraph("Черновик ___ [12] УТВЕРЖДАЮ МОДУЛЬНЫЙ ДОКУМЕНТ A/B")

    findings = check_draft_file(_docx_bytes(document))
    problems = [finding["problem"] for finding in findings]

    assert any("незаполненная линия" in problem for problem in problems)
    assert any("маркер примечания" in problem for problem in problems)
    assert any("служебный шаблонный текст" in problem for problem in problems)
    assert any("слэш" in problem for problem in problems)
    assert any("предупреждение" in problem for problem in problems)
    assert any("пояснительная записка" in problem for problem in problems)
